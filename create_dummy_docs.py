#!/usr/bin/env python3
"""Create two dummy legal contract DOCX files for testing."""

from docx import Document

# --- Document A: Casual legal contract ---
doc_a = Document()
doc_a.add_paragraph(
    "If the Buyer doesn't pay within 30 days, "
    "the Seller can basically cancel the whole deal and keep the deposit."
)
doc_a.add_paragraph(
    "Both sides agree to sort out any disagreements "
    "by talking it over before going to court or anything like that."
)
doc_a.save("doc_a.docx")
print("Created doc_a.docx (casual legal contract)")

# --- Document B: Formal legal contract (same subject) ---
doc_b = Document()
doc_b.add_paragraph(
    "In the event that the Buyer fails to remit payment within thirty (30) calendar days "
    "of the invoice date, the Seller shall have the right to terminate this Agreement "
    "and retain the deposit as liquidated damages."
)
doc_b.add_paragraph(
    "The Parties agree to submit any dispute arising out of or in connection with this Agreement "
    "to mandatory mediation prior to initiating any litigation or arbitration proceedings."
)
doc_b.save("doc_b.docx")
print("Created doc_b.docx (formal legal contract)")
