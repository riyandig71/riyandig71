#!/usr/bin/env python3
"""Create two 2-page A4 legal contract DOCX files for testing."""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.section import WD_ORIENT


def setup_a4(doc):
    """Configure A4 page size with standard margins."""
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)


def add_para(doc, text, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    p.paragraph_format.space_after = Pt(space_after)
    return p


# =====================================================================
# DOCUMENT A — Casual / informal legal contract
# =====================================================================
doc_a = Document()
setup_a4(doc_a)

add_para(doc_a, "SERVICE AGREEMENT", bold=True, size=14, space_after=12)
add_para(doc_a, "Between: PT Maju Bersama (us, the Company) and Mr. Budi Santoso (Client)", size=11)
add_para(doc_a, "Date: January 15, 2026", size=11, space_after=12)

add_para(doc_a, "1. WHAT THIS AGREEMENT IS ABOUT", bold=True, size=11)
add_para(doc_a, (
    "This agreement is basically about the consulting services that we're going to provide to the Client. "
    "We'll be helping with legal compliance stuff and regulatory advisory for the Client's business operations in Indonesia. "
    "The scope covers general legal consultation, contract drafting assistance, regulatory filing support, and compliance monitoring. "
    "If the Client needs extra services that aren't listed here, we can discuss and add them later with a separate addendum."
))

add_para(doc_a, "2. HOW LONG THIS LASTS", bold=True, size=11)
add_para(doc_a, (
    "This agreement starts on February 1, 2026 and goes for 12 months until January 31, 2027. "
    "If both sides are happy, we can renew it for another year by just agreeing in writing at least 30 days before it ends. "
    "Either party can end this agreement early by giving 60 days written notice to the other side. "
    "If someone terminates early, any fees already paid for services not yet delivered will be refunded on a pro-rata basis within 30 business days."
))

add_para(doc_a, "3. PAYMENT TERMS", bold=True, size=11)
add_para(doc_a, (
    "The Client will pay us IDR 50,000,000 per month for the services described above. "
    "Payment is due within 14 days after we send the invoice at the beginning of each month. "
    "If the Client doesn't pay on time, we'll charge a late fee of 1.5% per month on the outstanding amount. "
    "All payments should be made by bank transfer to our designated account. "
    "We reserve the right to pause our services if payment is overdue by more than 30 days."
))

add_para(doc_a, "4. WHAT WE PROMISE TO DO", bold=True, size=11)
add_para(doc_a, (
    "We promise to do our best to provide quality consulting services in line with professional standards. "
    "Our team will assign at least one senior consultant to handle the Client's account. "
    "We will respond to the Client's inquiries within 2 business days under normal circumstances. "
    "For urgent matters that affect the Client's legal standing, we will try to respond within 24 hours. "
    "We'll provide monthly summary reports of all activities and advice given during that period."
))

add_para(doc_a, "5. CONFIDENTIALITY", bold=True, size=11)
add_para(doc_a, (
    "Both sides agree to keep each other's confidential information secret and not share it with anyone else. "
    "This means all business information, financial data, trade secrets, client lists, and any documents shared during the engagement. "
    "This confidentiality obligation will continue for 3 years even after this agreement ends. "
    "If someone needs to disclose confidential info because of a court order or law, they need to tell the other party first whenever possible."
))

# Force a page break for page 2
doc_a.add_page_break()

add_para(doc_a, "6. WHO OWNS THE WORK", bold=True, size=11)
add_para(doc_a, (
    "All the work product and deliverables we create for the Client during this engagement will belong to the Client. "
    "This includes legal opinions, drafted contracts, compliance reports, and any advisory documents. "
    "However, our general know-how, methodologies, and tools that we developed independently remain our property. "
    "We can use general knowledge gained during this engagement for other clients as long as no confidential information is disclosed."
))

add_para(doc_a, "7. LIABILITY AND INDEMNIFICATION", bold=True, size=11)
add_para(doc_a, (
    "Our total liability under this agreement won't exceed the total fees paid by the Client in the last 6 months. "
    "We're not responsible for any indirect damages, lost profits, or consequential losses even if we were told they might happen. "
    "The Client agrees to indemnify us against any claims from third parties that arise from the Client's use of our advice in ways we didn't recommend. "
    "Neither party will be held responsible for delays or failures caused by events beyond their reasonable control like natural disasters, government actions, or pandemics."
))

add_para(doc_a, "8. DISPUTE RESOLUTION", bold=True, size=11)
add_para(doc_a, (
    "If there's a disagreement between us, we'll first try to sort it out through friendly discussion within 30 days. "
    "If that doesn't work, we agree to go to mediation using a mediator that both sides can agree on. "
    "If mediation also fails after 60 days, then either party can take the matter to the District Court of Central Jakarta. "
    "This agreement is governed by the laws of the Republic of Indonesia."
))

add_para(doc_a, "9. GENERAL STUFF", bold=True, size=11)
add_para(doc_a, (
    "This agreement is the entire deal between us and replaces any previous discussions or agreements about the same subject. "
    "Any changes to this agreement need to be in writing and signed by both parties to be valid. "
    "If any part of this agreement turns out to be unenforceable, the rest of it still stands. "
    "Neither side can transfer their rights or obligations under this agreement without the other's written consent. "
    "This agreement can be signed in multiple copies and each one counts as an original."
))

add_para(doc_a, "SIGNATURES", bold=True, size=11, space_after=20)
add_para(doc_a, "For PT Maju Bersama: ___________________     Date: ___________")
add_para(doc_a, "For Mr. Budi Santoso: ___________________    Date: ___________")

doc_a.save("doc_a.docx")
print("Created doc_a.docx (casual legal contract, 2 pages A4)")


# =====================================================================
# DOCUMENT B — Formal / benchmark legal contract (same subject)
# =====================================================================
doc_b = Document()
setup_a4(doc_b)

add_para(doc_b, "PROFESSIONAL SERVICES AGREEMENT", bold=True, size=14, space_after=12)
add_para(doc_b, "Between: PT Maju Bersama (hereinafter referred to as \"the Service Provider\") and Mr. Budi Santoso (hereinafter referred to as \"the Client\")", size=11)
add_para(doc_b, "Effective Date: January 15, 2026", size=11, space_after=12)

add_para(doc_b, "1. SCOPE OF SERVICES", bold=True, size=11)
add_para(doc_b, (
    "The Service Provider shall render professional consulting services to the Client encompassing legal compliance advisory and regulatory guidance pertaining to the Client's business operations within the jurisdiction of the Republic of Indonesia. "
    "The scope of services shall include, but not be limited to, general legal consultation, contract drafting and review, regulatory filing assistance, and ongoing compliance monitoring. "
    "Any services not expressly enumerated herein shall require a written addendum executed by both Parties prior to commencement."
))

add_para(doc_b, "2. TERM AND TERMINATION", bold=True, size=11)
add_para(doc_b, (
    "This Agreement shall commence on February 1, 2026 and shall remain in effect for a period of twelve (12) months, terminating on January 31, 2027, unless earlier terminated in accordance with the provisions herein. "
    "This Agreement may be renewed for successive twelve-month periods upon mutual written consent of both Parties, provided that such consent is communicated no fewer than thirty (30) calendar days prior to the expiration of the then-current term. "
    "Either Party may terminate this Agreement without cause by providing sixty (60) calendar days' prior written notice to the other Party. "
    "Upon early termination, any prepaid fees attributable to services not yet rendered shall be refunded to the Client on a pro-rata basis within thirty (30) business days of the effective date of termination."
))

add_para(doc_b, "3. FEES AND PAYMENT", bold=True, size=11)
add_para(doc_b, (
    "In consideration of the services rendered hereunder, the Client shall pay the Service Provider a monthly retainer fee of IDR 50,000,000 (fifty million Indonesian Rupiah). "
    "Payment shall be due and payable within fourteen (14) calendar days following the issuance of the Service Provider's invoice at the commencement of each calendar month. "
    "In the event of late payment, the Client shall be liable for a late payment penalty of 1.5% per month, calculated on the outstanding balance from the date such payment was due. "
    "All payments shall be remitted via bank transfer to the Service Provider's designated bank account as specified in the invoice. "
    "The Service Provider reserves the right to suspend the provision of services in the event that any payment remains outstanding for a period exceeding thirty (30) calendar days."
))

add_para(doc_b, "4. SERVICE PROVIDER'S OBLIGATIONS", bold=True, size=11)
add_para(doc_b, (
    "The Service Provider shall perform all services with the degree of skill, care, and diligence consistent with accepted professional standards and practices in the legal consulting industry. "
    "The Service Provider shall designate at least one (1) senior consultant to serve as the primary point of contact for the Client's account. "
    "The Service Provider shall use reasonable efforts to respond to the Client's inquiries within two (2) business days under ordinary circumstances. "
    "For matters of an urgent nature that may materially affect the Client's legal standing or regulatory compliance, the Service Provider shall endeavor to respond within twenty-four (24) hours. "
    "The Service Provider shall furnish monthly written reports summarizing all activities undertaken and advisory services rendered during the preceding calendar month."
))

add_para(doc_b, "5. CONFIDENTIALITY", bold=True, size=11)
add_para(doc_b, (
    "Each Party undertakes to maintain the strict confidentiality of all Confidential Information received from the other Party and shall not disclose such information to any third party without prior written consent. "
    "\"Confidential Information\" shall mean all business information, financial data, trade secrets, client lists, proprietary methodologies, and any documents or materials exchanged between the Parties in connection with this Agreement. "
    "The obligations of confidentiality set forth herein shall survive the termination or expiration of this Agreement for a period of three (3) years. "
    "Notwithstanding the foregoing, a Party may disclose Confidential Information to the extent required by applicable law, regulation, or court order, provided that the disclosing Party shall, to the extent legally permissible, provide the other Party with prompt written notice of such requirement prior to disclosure."
))

# Force a page break for page 2
doc_b.add_page_break()

add_para(doc_b, "6. INTELLECTUAL PROPERTY RIGHTS", bold=True, size=11)
add_para(doc_b, (
    "All work product, deliverables, and materials created by the Service Provider in the course of performing services under this Agreement shall constitute \"Work Product\" and shall be the exclusive property of the Client upon full payment of all applicable fees. "
    "Work Product shall include, without limitation, legal opinions, drafted agreements, compliance reports, regulatory submissions, and advisory memoranda. "
    "Notwithstanding the foregoing, the Service Provider shall retain all rights, title, and interest in and to its pre-existing intellectual property, proprietary methodologies, tools, and general professional know-how. "
    "The Service Provider may utilize general skills, knowledge, and experience acquired during the performance of this Agreement for the benefit of other clients, provided that no Confidential Information of the Client is disclosed or utilized in connection therewith."
))

add_para(doc_b, "7. LIMITATION OF LIABILITY AND INDEMNIFICATION", bold=True, size=11)
add_para(doc_b, (
    "The aggregate liability of the Service Provider under or in connection with this Agreement shall not exceed the total fees actually paid by the Client during the six (6) month period immediately preceding the event giving rise to such liability. "
    "In no event shall either Party be liable to the other for any indirect, incidental, special, consequential, or punitive damages, including but not limited to loss of profits, loss of revenue, or loss of business opportunity, regardless of whether such Party has been advised of the possibility of such damages. "
    "The Client shall indemnify, defend, and hold harmless the Service Provider from and against any and all claims, liabilities, damages, and expenses arising out of or relating to the Client's implementation or utilization of the Service Provider's advice in a manner inconsistent with the Service Provider's express recommendations. "
    "Neither Party shall be liable for any delay or failure to perform its obligations under this Agreement to the extent that such delay or failure is attributable to Force Majeure events, including but not limited to natural disasters, acts of government, epidemics, or pandemics."
))

add_para(doc_b, "8. DISPUTE RESOLUTION AND GOVERNING LAW", bold=True, size=11)
add_para(doc_b, (
    "In the event of any dispute, controversy, or claim arising out of or in connection with this Agreement, the Parties shall first attempt to resolve the matter through good-faith negotiation for a period of thirty (30) calendar days. "
    "If the dispute cannot be resolved through negotiation, the Parties agree to submit the matter to mediation administered by a mutually agreed-upon mediator in Jakarta, Indonesia. "
    "In the event that mediation fails to resolve the dispute within sixty (60) calendar days of its commencement, either Party may submit the dispute to the exclusive jurisdiction of the District Court of Central Jakarta. "
    "This Agreement shall be governed by and construed in accordance with the laws of the Republic of Indonesia."
))

add_para(doc_b, "9. GENERAL PROVISIONS", bold=True, size=11)
add_para(doc_b, (
    "This Agreement constitutes the entire agreement between the Parties with respect to the subject matter hereof and supersedes all prior negotiations, representations, warranties, commitments, offers, and agreements, whether written or oral. "
    "No amendment, modification, or waiver of any provision of this Agreement shall be effective unless made in writing and duly executed by authorized representatives of both Parties. "
    "If any provision of this Agreement is held to be invalid, illegal, or unenforceable, the remaining provisions shall continue in full force and effect. "
    "Neither Party may assign, transfer, or delegate any of its rights or obligations under this Agreement without the prior written consent of the other Party. "
    "This Agreement may be executed in counterparts, each of which shall be deemed an original and all of which together shall constitute one and the same instrument."
))

add_para(doc_b, "IN WITNESS WHEREOF", bold=True, size=11, space_after=20)
add_para(doc_b, "For and on behalf of PT Maju Bersama: ___________________     Date: ___________")
add_para(doc_b, "For and on behalf of Mr. Budi Santoso: ___________________    Date: ___________")

doc_b.save("doc_b.docx")
print("Created doc_b.docx (formal legal contract, 2 pages A4)")
